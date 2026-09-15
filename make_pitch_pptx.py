# -*- coding: utf-8 -*-
"""华润三九杯 · 睡眠中控 4 分钟路演稿。运行：python make_pitch_pptx.py"""
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

OUT = r"d:\pro222\sleep-hub\睡眠中控-路演4分钟.pptx"

W, H = Inches(13.333), Inches(7.5)
BG = RGBColor(0x12, 0x11, 0x0F)
CREAM = RGBColor(0xF3, 0xEC, 0xE2)
MUTED = RGBColor(0xA8, 0x9F, 0x93)
DAWN = RGBColor(0xE0, 0xA8, 0x72)
LINE = RGBColor(0x3A, 0x34, 0x2C)
FONT = "微软雅黑"


def _ea(run, name):
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    xml = (
        '<a:ea xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
        f' typeface="{name}"/>'
    )
    node = etree.fromstring(xml)
    if ea is None:
        rPr.append(node)
    else:
        ea.set("typeface", name)


def paint(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG


def box(slide, l, t, w, h):
    s = slide.shapes.add_textbox(l, t, w, h)
    tf = s.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    return tf


def run(p, text, size, color, bold=False, font=FONT):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = font
    _ea(r, font)
    return r


def para(tf, text, size, color, bold=False, align=PP_ALIGN.LEFT, space_before=0, space_after=6):
    p = tf.paragraphs[0] if not tf.paragraphs[0].text and not tf.paragraphs[0].runs else tf.add_paragraph()
    if tf.paragraphs[0] is p and not p.runs:
        pass
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    if p.runs:
        p.clear()
    run(p, text, size, color, bold)
    return p


def first(tf, text, size, color, bold=False, align=PP_ALIGN.LEFT, after=8):
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(after)
    run(p, text, size, color, bold)
    return p


def add_line(slide, l, t, w):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, Emu(12700))
    sh.fill.solid()
    sh.fill.fore_color.rgb = DAWN
    sh.line.fill.background()


def card(slide, l, t, w, h):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    sh.adjustments[0] = 0.08
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor(0x1C, 0x1A, 0x17)
    sh.line.color.rgb = LINE
    sh.line.width = Pt(1)
    return sh


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def footer(slide, n, total=8):
    tf = box(slide, Inches(0.7), Inches(7.05), Inches(10.5), Inches(0.3))
    first(tf, "睡眠中控", 11, MUTED, after=0)
    tf2 = box(slide, Inches(11.6), Inches(7.05), Inches(1.0), Inches(0.3))
    first(tf2, f"{n} / {total}", 11, MUTED, align=PP_ALIGN.RIGHT, after=0)


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]

    # 1 开场：人被叫起来，但房间没参与
    s = prs.slides.add_slide(blank)
    paint(s)
    add_line(s, Inches(0.85), Inches(1.85), Inches(1.4))
    tf = box(s, Inches(0.85), Inches(2.1), Inches(11.5), Inches(0.95))
    first(tf, "睡眠中控", 48, CREAM, True, after=0)
    tf = box(s, Inches(0.85), Inches(3.2), Inches(11.5), Inches(0.55))
    first(tf, "指定采集，指定叫醒", 26, DAWN, after=0)
    tf = box(s, Inches(0.85), Inches(4.0), Inches(11.5), Inches(0.55))
    first(tf, "窗口到了，房间里要真的发生事。", 18, MUTED, after=0)
    notes(
        s,
        "0:00–0:25  宿舍里很多人戴着手环。早上被手机硬拽起来，App 里有一串分数，灯和铃却各管各的。"
        "该叫的窗口到了，房间常常什么都没发生。我们做的不是再测一次，是让介入发生。",
    )

    # 2 冲突
    s = prs.slides.add_slide(blank)
    paint(s)
    tf = box(s, Inches(0.85), Inches(0.7), Inches(11.5), Inches(0.4))
    first(tf, "昨晚被测完，今早房间里没动", 14, DAWN, True, after=0)
    beats = [
        ("手环", "记下分数，窗口里写不出「现在浅睡」。"),
        ("灯和手机", "到点就亮、到点就响，不管人在哪一段。"),
        ("焊死的套装", "换一个采集源，叫醒步骤就要重做。"),
    ]
    for i, (h, b) in enumerate(beats):
        y = Inches(1.4 + i * 1.55)
        card(s, Inches(0.85), y, Inches(11.6), Inches(1.35))
        tf = box(s, Inches(1.2), y + Inches(0.22), Inches(10.9), Inches(0.4))
        first(tf, h, 20, DAWN, True, after=0)
        tf = box(s, Inches(1.2), y + Inches(0.68), Inches(10.9), Inches(0.45))
        first(tf, b, 18, CREAM, after=0)
    footer(s, 2)
    notes(
        s,
        "0:25–0:50  缺的不是分数。缺的是浅睡到了，灯、铃、手机会不会动。不要点名竞品。",
    )

    # 3 主办给了方向
    s = prs.slides.add_slide(blank)
    paint(s)
    tf = box(s, Inches(0.85), Inches(1.55), Inches(11.5), Inches(0.4))
    first(tf, "主办已经把方向写下了", 14, DAWN, True, after=0)
    tf = box(s, Inches(0.85), Inches(2.1), Inches(11.5), Inches(0.9))
    first(tf, "测得更准  →  介入得更对", 34, CREAM, True, after=0)
    tf = box(s, Inches(0.85), Inches(3.25), Inches(11.5), Inches(1.1))
    first(tf, "思考入口不是标准答案。我们不新做助眠 App，把介入写成能跑的一层。", 18, MUTED, after=0)
    footer(s, 3)
    notes(
        s,
        "0:50–1:15  点大图原词就够。下一句落到中控：采集谁来写可以换，叫醒怎么执行可以换，中间不变。",
    )

    # 4 怎么跑
    s = prs.slides.add_slide(blank)
    paint(s)
    tf = box(s, Inches(0.85), Inches(0.7), Inches(11.5), Inches(0.4))
    first(tf, "中控只认今夜三件事", 14, DAWN, True, after=0)
    tf = box(s, Inches(0.85), Inches(1.2), Inches(11.5), Inches(0.8))
    first(tf, "在不在  ·  动没动  ·  到了哪一段", 28, CREAM, True, after=0)
    cols = [
        ("采集", "点选、HTTP，以后手环或雷达。谁写都可以。"),
        ("配方", "等浅睡，再黎明、响铃、推手机。夜里不跑模型。"),
        ("执行", "屏幕或灯、电脑铃、手机。灯到了只改一个地址。"),
    ]
    for i, (h, b) in enumerate(cols):
        x = Inches(0.85 + i * 4.0)
        card(s, x, Inches(2.35), Inches(3.7), Inches(3.55))
        tf = box(s, x + Inches(0.28), Inches(2.6), Inches(3.15), Inches(0.5))
        first(tf, h, 20, DAWN, True, after=0)
        tf = box(s, x + Inches(0.28), Inches(3.25), Inches(3.15), Inches(2.2))
        first(tf, b, 16, CREAM, after=0)
    footer(s, 4)
    notes(
        s,
        "1:15–1:40  三列指过去。口述可以编译成配方，守夜只跑契约。准备切浏览器。",
    )

    # 5 高潮：演示
    s = prs.slides.add_slide(blank)
    paint(s)
    add_line(s, Inches(0.85), Inches(2.2), Inches(1.4))
    tf = box(s, Inches(0.85), Inches(2.45), Inches(11.5), Inches(0.85))
    first(tf, "今夜这一轮", 40, CREAM, True, after=0)
    tf = box(s, Inches(0.85), Inches(3.5), Inches(11.5), Inches(0.7))
    first(tf, "指定浅睡，然后亮、响、震。", 22, MUTED, after=0)
    footer(s, 5)
    notes(
        s,
        "1:40–3:10  立刻切中控。「采集是我指定的。」一键三幕。亮、响、震。「叫醒也是我指定的。」"
        "有第二台就 POST 一条浅睡，步骤不改：「换的是谁来写事实，不是换产品。」失败放备放录像。",
    )

    # 6 明天怎么续
    s = prs.slides.add_slide(blank)
    paint(s)
    tf = box(s, Inches(0.85), Inches(1.7), Inches(11.5), Inches(0.45))
    first(tf, "换源不换方案", 14, DAWN, True, after=0)
    tf = box(s, Inches(0.85), Inches(2.2), Inches(11.5), Inches(0.85))
    first(tf, "产品不是某一颗传感器。", 32, CREAM, True, after=0)
    tf = box(s, Inches(0.85), Inches(3.3), Inches(11.5), Inches(1.2))
    first(tf, "雷达、灯到了，只接线、改一步地址。单通道脑电只证明戴上了，不当分期。", 18, MUTED, after=0)
    footer(s, 6)
    notes(
        s,
        "3:10–3:35  不报型号，不解释快递。被问：模组已下单，周末未发货，口子留好了。",
    )

    # 7 收束
    s = prs.slides.add_slide(blank)
    paint(s)
    add_line(s, Inches(0.85), Inches(2.15), Inches(1.4))
    tf = box(s, Inches(0.85), Inches(2.4), Inches(11.5), Inches(1.0))
    first(tf, "采集是我指定的。叫醒是我指定的。", 28, CREAM, True, after=0)
    tf = box(s, Inches(0.85), Inches(3.55), Inches(11.5), Inches(0.5))
    first(tf, "现场已经执行完一轮。请提问。", 18, MUTED, after=0)
    footer(s, 7)
    notes(s, "3:35–4:00  停。不要加准确率、融资。")

    # 8 问答：一句完整话，不投影
    s = prs.slides.add_slide(blank)
    paint(s)
    tf = box(s, Inches(0.85), Inches(0.5), Inches(11.5), Inches(0.35))
    first(tf, "附录 · 不投影", 14, DAWN, True, after=0)
    qa = [
        ("没硬件？", "产品是中控。指定采集和叫醒，不靠快递。"),
        ("AI？", "编译口述时用，夜里配方自己跑。"),
        ("准确率？", "不宣称分期。我们交窗口到了会执行。"),
        ("和手环闹钟？", "手环多是夜后报告；我们窗口内执行，源可换。"),
        ("雷达 / 脑电？", "都可插。雷达写在不在；脑电只写戴上没有。"),
        ("三九能用吗？", "家庭侧先证明房间里会动。"),
    ]
    for i, (q, a) in enumerate(qa):
        col, row = i % 2, i // 2
        x = Inches(0.85 + col * 6.15)
        y = Inches(1.05 + row * 1.75)
        card(s, x, y, Inches(5.9), Inches(1.55))
        tf = box(s, x + Inches(0.28), y + Inches(0.18), Inches(5.35), Inches(0.4))
        first(tf, q, 16, DAWN, True, after=0)
        tf = box(s, x + Inches(0.28), y + Inches(0.65), Inches(5.35), Inches(0.7))
        first(tf, a, 15, CREAM, after=0)
    footer(s, 8)
    notes(s, "每问一句完整话。不谈型号、融资。")

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
